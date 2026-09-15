import pathlib
ROOT = pathlib.Path(__file__).resolve().parents[1]
RAW  = ROOT / "data" / "raw"
DER  = ROOT / "data" / "derived"
OUT  = ROOT / "output"
import math
import pandas as pd
from openpyxl import load_workbook

U = str(RAW) + "/"
MM = {'Jan':1,'Shk':2,'Mar':3,'Pri':4,'Maj':5,'Qer':6,'Korr':7,'Gsh':8,'Gush':8,
      'Sht':9,'Tet':10,'Nen':11,'Nën':11,'Dhj':12}

def bsh(fname):
    """Read a BSH time-series export -> {code: {(year,month): value}}"""
    ws = load_workbook(U+fname, read_only=True)[load_workbook(U+fname, read_only=True).sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[7]
    out = {}
    for r in rows[8:]:
        if r[0] is None:
            continue
        d = {}
        for i in range(2, len(r)):
            if hdr[i] is None or r[i] is None:
                continue
            p = str(hdr[i]).strip().split()
            if p[0] in MM:
                d[(int(p[1]), MM[p[0]])] = r[i]
        out[str(r[0])] = d
    return out

QS = [(y, m) for y in range(2016, 2027) for m in (3, 6, 9, 12)]
QS = [q for q in QS if not (q[0] == 2026 and q[1] > 6)]
def lab(q): return f"{q[0]}Q{q[1]//3}"

df = pd.DataFrame(index=[lab(q) for q in QS])

# ---- 1. NPL, kapitali -------------------------------------------------
fsi = bsh("boa_financial_soundness_indicators.xlsx")
df["npl"]      = [fsi["2.6"].get(q)   for q in QS]
df["car"]      = [fsi["1.1"].get(q)   for q in QS]
df["npl_kap"]  = [fsi["1.4.1"].get(q) for q in QS]
df["roa"]      = [fsi["2.5"].get(q)   for q in QS]

# ---- 2. Kredia sipas qëllimit: hipoteka + pasuri të paluajtshme biznes
q9 = bsh("boa_loans_by_purpose_currency.xlsx")
mort = ["1.3.1.4", "1.3.2.4", "1.3.3.4", "1.3.4.4"]          # individë, blerje banesash
bre  = ["1.2.1.5", "1.2.2.5", "1.2.3.5", "1.2.4.5"]          # biznes privat, pasuri të paluajtshme
df["kredi_tot"] = [q9["1"].get(q) for q in QS]
df["hipoteka"]  = [sum(q9[c].get(q) or 0 for c in mort) for q in QS]
df["pp_biznes"] = [sum(q9[c].get(q) or 0 for c in bre)  for q in QS]

# ---- 3. Ekspozimi i biznesit ndaj ndërtimit+PP (sipas rrethit) --------
ws = load_workbook(U+"boa_loans_by_district_activity.xlsx", read_only=True)["Agregate"]
rows = list(ws.iter_rows(values_only=True)); hdr = rows[7]
cols = {}
for i in range(2, len(hdr)):
    if hdr[i] is None: continue
    t = str(hdr[i]).replace("T ", "T").strip().split()      # 'T III 2016'
    roman = {"TI":1,"TII":2,"TIII":3,"TIV":4}
    key = t[0]
    cols[(int(t[-1]), roman["T"+t[-2] if len(t)==3 else t[0]])] = i
d3 = {}
for r in rows[8:]:
    if r[0] is None: continue
    d3[str(r[0])] = {k: r[i] for k, i in cols.items()}
dist = ["1.1","1.2","1.3","1.4","1.5","1.6","1.7","1.8","1.9"]
def exp_q(q):
    key = (q[0], q[1]//3)
    if key not in d3["1"] or d3["1"][key] in (None, 0): return None
    num = sum((d3[k+".6"][key] or 0) + (d3[k+".12"][key] or 0) for k in dist)
    return 100*num/d3["1"][key]
df["eksp_ndertim"] = [exp_q(q) for q in QS]

# ---- 4. PBB reale, rritje vjetore ------------------------------------
ws = load_workbook(U+"instat_gdp_quarterly_production.xlsx", read_only=True)["Rritja reale tremujore"]
rows = list(ws.iter_rows(values_only=True))
gdp, cons, year = {}, {}, None
for r in rows:
    a = str(r[0]).strip() if r[0] else ""
    if a[:4].isdigit() and len(a) >= 4: year = int(a[:4])
    b = str(r[1]).strip() if r[1] else ""
    if year and b in ("T1","T2","T3","T4") and r[15] is not None:
        gdp[(year, int(b[1])*3)]  = r[15]
        cons[(year, int(b[1])*3)] = r[5]
df["pbb"]      = [gdp.get(q)  for q in QS]
df["ndertim"]  = [cons.get(q) for q in QS]

# ---- 5. Norma bazë (repo), forward-fill ------------------------------
ws = load_workbook(U+"boa_policy_rates.xlsx", read_only=True)["Normat_Rates"]
EN = {'Jan':1,'Feb':2,'Mar':3,'Apr':4,'May':5,'Jun':6,'Jul':7,'Aug':8,'Sep':9,'Oct':10,'Nov':11,'Dec':12}
ev, year = [], None
for r in ws.iter_rows(values_only=True):
    a = str(r[0]).strip() if r[0] else ""
    if a.isdigit() and len(a) == 4: year = int(a); continue
    if r[2] is not None and "-" in a:
        d, mo = a.split("-")[0], a.split("-")[1]
        if mo in EN: ev.append((year, EN[mo], int(d), float(r[2])))
ev.sort()
def repo(q):
    v = None
    for (y, m, d, x) in ev:
        if (y, m) <= q: v = x
    return v
df["repo"] = [repo(q) for q in QS]

# ---- 6. Inflacioni (ndryshim vjetor, mesatare tremujore) -------------
infl = {}
p16 = {1:1.5, 2:0.2, 3:0.3, 4:0.3, 5:0.7, 6:1.2, 7:1.9, 8:2.0, 9:1.8, 10:1.5, 11:1.9, 12:2.2}
for m, v in p16.items(): infl[(2016, m)] = v
ws = load_workbook(U+"instat_cpi_annual_change.xlsx", read_only=True)["Sheet1"]
rows = list(ws.iter_rows(values_only=True)); hdr = rows[3]
for r in rows[4:]:
    if r[0] is not None and str(r[0]).strip() == "000000":
        for i in range(2, len(hdr)):
            if hdr[i] is None or r[i] is None: continue
            t = str(hdr[i]).split()[0]
            if "-" in t:
                mo, yr = t.split("-")
                if mo.isdigit(): infl[(2000+int(yr), int(mo))] = float(r[i])
        break
a8 = bsh("boa_cpi_2022_2026.xlsx")                     # mbulon 2026
for (y, m), v in a8["B"].items():
    if (y, m) not in infl: infl[(y, m)] = v
def infl_q(q):
    ms = [infl.get((q[0], q[1]-2)), infl.get((q[0], q[1]-1)), infl.get((q[0], q[1]))]
    ms = [x for x in ms if x is not None]
    return sum(ms)/len(ms) if ms else None
df["inflacion"] = [infl_q(q) for q in QS]

# ---- 7. HPI gjashtëmujor -> tremujor ---------------------------------
nat = {"2020H1":1.400,"2020H2":1.487,"2021H1":1.495,"2021H2":1.621,"2022H1":2.081,
       "2022H2":1.892,"2023H1":2.111,"2023H2":1.950,"2024H1":2.280,"2024H2":2.819,
       "2025H1":3.231,"2025H2":3.609}
tir = {"2020H1":1.434,"2020H2":1.467,"2021H1":1.518,"2021H2":1.647,"2022H1":2.117,
       "2022H2":1.876,"2023H1":2.125,"2023H2":1.910,"2024H1":2.359,"2024H2":2.977,
       "2025H1":3.129,"2025H2":3.129}
def to_q(h):
    s = {}
    for k, v in h.items():
        y = int(k[:4]); s[(y, 2 if k[-1] == "1" else 4)] = v      # H1->Q2, H2->Q4
    out = {}
    ks = sorted(s)
    for i, k in enumerate(ks):
        y, qq = k
        out[(y, qq*3)] = s[k]
        if i > 0:
            py, pq = ks[i-1]
            prev = s[ks[i-1]]
            mid = (prev + s[k]) / 2
            my, mq = (y, 1) if qq == 2 else (y, 3)
            out[(my, mq*3)] = mid
    return out
df["hpi_vendi"]  = [to_q(nat).get(q) for q in QS]
df["hpi_tirane"] = [to_q(tir).get(q) for q in QS]

# ---- 8. Indeksi i kushtimit në ndërtim -------------------------------
ws = load_workbook(U+"instat_construction_cost_index.xlsx", read_only=True)["sheet 1"]
rows = list(ws.iter_rows(values_only=True)); hdr = rows[4]
R = {"I":1,"II":2,"III":3,"IV":4}
kk = {}
for r in rows:
    if r[0] and "TOTALI" in str(r[0]):
        for i in range(1, len(hdr)):
            if hdr[i] is None or r[i] is None: continue
            t = str(hdr[i]).strip().split("-")
            if t[0] in R: kk[(2000+int(t[1]), R[t[0]]*3)] = float(r[i])
        break
df["kosto_ndertim"] = [kk.get(q) for q in QS]

# ---- 9. Transformimet ------------------------------------------------
df["npl_logit"]  = df["npl"].apply(lambda x: math.log(x/(100-x)) if pd.notna(x) else None)
df["stok_npl"]   = df["npl"]/100 * df["kredi_tot"]
df["ln_stok"]    = df["stok_npl"].apply(lambda x: math.log(x) if pd.notna(x) else None)
df["hip_share"]  = 100*df["hipoteka"]/df["kredi_tot"]
df["ppb_share"]  = 100*df["pp_biznes"]/df["kredi_tot"]
df["eksp_total"] = df["hip_share"] + df["ppb_share"]
df["repo_real"]  = df["repo"] - df["inflacion"]
df["ln_hpi"]     = df["hpi_vendi"].apply(lambda x: math.log(x) if pd.notna(x) else None)
df["ln_hpi_tr"]  = df["hpi_tirane"].apply(lambda x: math.log(x) if pd.notna(x) else None)
df["rritje_kredi"] = 100*(df["kredi_tot"]/df["kredi_tot"].shift(4) - 1)
df["d_covid"]    = [1 if q in [(2020,6),(2020,9),(2020,12),(2021,3)] else 0 for q in QS]
df["d_ltv"]      = [1 if (q[0],q[1]) >= (2025,9) else 0 for q in QS]

df.index.name = "tremujori"
df = df.round(4)
df.to_csv(str(DER / "dataset_quarterly_2016Q1_2026Q2.csv"))
print(df.to_string())
